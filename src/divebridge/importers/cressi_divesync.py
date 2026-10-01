"""Importer for the XLSX export of the DiveSync app (Cressi Da Vinci, XS Scuba Skiff, ...).

Workbook layout (DiveSync "DiveLogExport" v1):
  DiveLog            one row per dive, imperial units (ft, °F, psi)
  DiveProfile        one row per sample, metric units (m, °C), DiveTime in seconds
  TankData           four rows per dive (tank 1-4), often empty
  _DiveSync_Metadata hidden sheet: App=DiveSync, FileType=DiveLogExport, Version=1
"""

from __future__ import annotations

import io
import warnings
import zipfile
from collections import defaultdict
from datetime import datetime
from typing import Any

import openpyxl

from ..model import Dive, DiveComputer, Gas, Sample, Site, Source, Tank
from ..units import f_to_c, ft_to_m, psi_to_bar, to_float, to_int, u16_or_none, u16_signed
from .base import ImportError_

FORMAT = "cressi_divesync"

# DiveSync reports the hardware platform name; map it to the marketed product.
DEVICE_NAMES = {
    "SKIFF": "Da Vinci",
}
MANUFACTURER = "Cressi"

DATE_FORMATS = (
    "%m/%d/%Y %H:%M:%S",
    "%m/%d/%Y %I:%M:%S %p",
    "%d/%m/%Y %H:%M:%S",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%dT%H:%M:%S",
    "%d.%m.%Y %H:%M:%S",
)


def _sheet_rows(ws) -> list[dict[str, Any]]:
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        return []
    keys = [str(h).strip() if h is not None else "" for h in header]
    out = []
    for row in rows:
        if row is None or all(v is None or v == "" for v in row):
            continue
        out.append({k: v for k, v in zip(keys, row) if k})
    return out


def _parse_datetime(value: Any, date_format_hint: int | None) -> datetime:
    if isinstance(value, datetime):
        return value
    s = str(value).strip()
    formats = list(DATE_FORMATS)
    if date_format_hint == 1:  # assume 1 = DMY; untested – try it before the MDY default
        formats.remove("%d/%m/%Y %H:%M:%S")
        formats.insert(0, "%d/%m/%Y %H:%M:%S")
    for fmt in formats:
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    raise ImportError_(f"unrecognised date/time: {s!r}")


class CressiDiveSyncImporter:
    name = FORMAT
    description = "Cressi DiveSync app – Excel export (DiveLog/DiveProfile/TankData)"
    extensions = (".xlsx",)

    def __init__(self, manufacturer: str = MANUFACTURER, device_names: dict[str, str] | None = None):
        self.manufacturer = manufacturer
        self.device_names = device_names or DEVICE_NAMES

    # -- detection ---------------------------------------------------------
    def can_handle(self, filename: str, data: bytes) -> bool:
        if not data.startswith(b"PK"):
            return False
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                wb = zf.read("xl/workbook.xml").decode("utf-8", "replace")
        except (zipfile.BadZipFile, KeyError):
            return False
        return "DiveLog" in wb and "DiveProfile" in wb

    # -- parsing -----------------------------------------------------------
    def parse(self, filename: str, data: bytes) -> list[Dive]:
        try:
            # not read_only: DiveSync writes a bogus <dimension ref="A1"/> which read_only mode trusts
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")  # "Workbook contains no default style"
                wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        except Exception as e:  # noqa: BLE001
            raise ImportError_(f"{filename}: not a readable xlsx ({e})") from e
        names = set(wb.sheetnames)
        for required in ("DiveLog", "DiveProfile"):
            if required not in names:
                raise ImportError_(f"{filename}: sheet {required!r} missing")

        meta: dict[str, Any] = {}
        if "_DiveSync_Metadata" in names:
            for row in wb["_DiveSync_Metadata"].iter_rows(values_only=True):
                if row and row[0] is not None:
                    meta[str(row[0])] = row[1] if len(row) > 1 else None

        logs = _sheet_rows(wb["DiveLog"])
        profiles: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for r in _sheet_rows(wb["DiveProfile"]):
            profiles[str(to_int(r.get("DiveID")))].append(r)
        tanks: dict[str, list[dict[str, Any]]] = defaultdict(list)
        if "TankData" in names:
            for r in _sheet_rows(wb["TankData"]):
                tanks[str(to_int(r.get("DiveID")))].append(r)
        wb.close()

        dives: list[Dive] = []
        for row in logs:
            if to_int(row.get("IsDelete")) == 1:
                continue
            dive_id = str(to_int(row.get("DiveID")))
            dives.append(self._dive(filename, row, profiles.get(dive_id, []), tanks.get(dive_id, []), meta))
        dives.sort(key=lambda d: d.start)
        return dives

    def _dive(self, filename: str, row: dict[str, Any], profile: list[dict[str, Any]],
              tank_rows: list[dict[str, Any]], meta: dict[str, Any]) -> Dive:
        dive_id = str(to_int(row.get("DiveID")))
        start = _parse_datetime(row.get("DiveStartLocalTime"), to_int(row.get("DateFormat")))

        # gases: 7 mix slots, the one in use is StartingMixIdx (1-based)
        mixes: list[Gas] = []
        for i in range(1, 8):
            o2 = to_float(row.get(f"Mix{i}Fo2Percent"))
            he = to_float(row.get(f"Mix{i}FHePercent")) or 0.0
            mixes.append(Gas(o2=o2 if o2 is not None else 21.0, he=he))
        start_idx = to_int(row.get("StartingMixIdx")) or 1
        used_indices: list[int] = [start_idx]
        for r in profile:
            idx = to_int(r.get("CurrentUsedMixIdx"))
            if idx and idx not in used_indices:
                used_indices.append(idx)
        valid_indices = [i for i in used_indices if 1 <= i <= 7]
        gases = [mixes[i - 1] for i in valid_indices]
        gas_pos = {idx: pos for pos, idx in enumerate(valid_indices)}  # sentinels (e.g. 255) -> no gas index

        samples: list[Sample] = []
        for r in profile:
            depth = to_float(r.get("DepthMeters"))
            t = to_float(r.get("DiveTime"))
            if depth is None or t is None:
                continue
            deco_depth_ft = to_float(r.get("DecoStopDepthFT"))
            pressure_psi = u16_or_none(r.get("TankPSI"))
            ppo2 = to_float(r.get("PpO2Barx100"))
            idx = to_int(r.get("CurrentUsedMixIdx"))
            alarms = tuple(
                str(r[k]).strip() for k in ("AlarmID1", "AlarmID2") if r.get(k) not in (None, "")
            )
            samples.append(
                Sample(
                    t_s=t,
                    depth_m=depth,
                    temp_c=to_float(r.get("TemperatureC")),
                    ndl_min=to_float(r.get("NdlMin")),
                    deco_depth_m=ft_to_m(deco_depth_ft) if deco_depth_ft else None,
                    deco_time_min=to_float(r.get("DecoTime")) or None,
                    pressure_bar=psi_to_bar(pressure_psi) if pressure_psi is not None else None,
                    ppo2_bar=ppo2 / 100.0 if ppo2 is not None else None,
                    gas_index=gas_pos.get(idx) if idx else None,
                    speed_m_min=(lambda v: ft_to_m(v) if v is not None else None)(u16_signed(r.get("SpeedFpm"))),
                    alarms=alarms,
                )
            )
        samples.sort(key=lambda s: s.t_s)

        tank_list: list[Tank] = []
        start_psi = u16_or_none(row.get("StartDiveTankPressurePSI"))
        end_psi = u16_or_none(row.get("EndDiveTankPressurePSI"))
        for tr in sorted(tank_rows, key=lambda x: to_int(x.get("TankNo")) or 0):
            vol = to_float(tr.get("CylinderSize"))
            wp = to_float(tr.get("WorkingPressure"))
            sp = u16_or_none(tr.get("StartPressure"))
            ep = u16_or_none(tr.get("EndPressure"))
            if vol is None and sp is None and ep is None:
                continue
            unit = to_int(tr.get("TankUnit")) or 0  # 0 = imperial (cuft/psi) assumed, 1 = metric (l/bar)
            if unit == 1 or vol is None:
                volume_l = vol
            else:
                # imperial "size" is gas capacity in cuft at working pressure (AL80 = 80 cuft @ 3000 psi
                # ≈ 11.1 l water volume); without a working pressure assume the usual 3000 psi
                wp_psi = wp if wp else 3000.0
                volume_l = vol * 28.3168 / (wp_psi / 14.696)
            tank_list.append(
                Tank(
                    index=to_int(tr.get("TankNo")) or len(tank_list) + 1,
                    volume_l=round(volume_l, 2) if volume_l is not None else None,
                    working_pressure_bar=wp if unit == 1 else (psi_to_bar(wp) if wp is not None else None),
                    start_bar=sp if unit == 1 else (psi_to_bar(sp) if sp is not None else None),
                    end_bar=ep if unit == 1 else (psi_to_bar(ep) if ep is not None else None),
                    gas=gases[0] if gases else None,
                )
            )
        if not tank_list and (start_psi is not None or end_psi is not None):
            tank_list.append(
                Tank(
                    index=1,
                    start_bar=psi_to_bar(start_psi) if start_psi is not None else None,
                    end_bar=psi_to_bar(end_psi) if end_psi is not None else None,
                    gas=gases[0] if gases else None,
                )
            )

        max_ft = to_float(row.get("MaxDepthFT"))
        avg_ft = to_float(row.get("AvgDepthFT"))
        min_f = to_float(row.get("MinTemperatureF"))
        max_f = to_float(row.get("MaxTemperatureF"))
        temps = [s.temp_c for s in samples if s.temp_c is not None]

        max_depth = ft_to_m(max_ft) if max_ft is not None else max((s.depth_m for s in samples), default=0.0)
        duration = to_int(row.get("TotalDiveTime"))
        if duration is None:
            duration = int(samples[-1].t_s) if samples else 0

        device = str(row.get("DeviceName") or "").strip() or "Unknown"
        serial = str(row.get("SerialNo") or "").strip() or None
        computer = DiveComputer(
            manufacturer=self.manufacturer,
            model=self.device_names.get(device.upper(), device),
            serial=serial,
            firmware=None,
        )
        site_name = str(row.get("DiveSiteName") or "").strip() or None
        memo = str(row.get("Memo") or "").strip() or None

        return Dive(
            start=start,
            duration_s=duration,
            max_depth_m=round(max_depth, 2),
            avg_depth_m=round(ft_to_m(avg_ft), 2) if avg_ft is not None else None,
            number=to_int(row.get("LifeTimeDiveNo")),
            water_temp_min_c=round(f_to_c(min_f), 1) if min_f is not None else (round(min(temps), 1) if temps else None),
            water_temp_max_c=round(f_to_c(max_f), 1) if max_f is not None else (round(max(temps), 1) if temps else None),
            surface_interval_s=to_int(row.get("SurfTime")),
            site=Site(name=site_name) if site_name else None,
            computer=computer,
            gases=gases,
            tanks=tank_list,
            samples=samples,
            notes=memo,
            deco=to_int(row.get("IsDecoDive")) == 1,
            gf_low=to_int(row.get("GfLowPercent")),
            gf_high=to_int(row.get("GfHighPercent")),
            source=Source(format=FORMAT, filename=filename, dive_id=dive_id),
            extra={
                "device_platform": device,
                "model_id": row.get("ModelID"),
                "dive_mode": to_int(row.get("DiveMode")),
                "sampling_s": to_int(row.get("SamplingTime")),
                "water_density": to_int(row.get("WaterDensity")),  # DiveSync setting; 0 observed with "salt" in the app
                "altitude_level": to_int(row.get("AltitudeLevel")),
                "export_version": meta.get("Version"),
                "exported_at": meta.get("ExportedAt"),
            },
        )
